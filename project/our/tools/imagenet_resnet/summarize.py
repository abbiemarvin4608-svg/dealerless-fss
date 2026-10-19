#!/usr/bin/env python3
"""Collect completed and incomplete ImageNet accuracy and protocol evidence."""
import argparse
import hashlib
import json
import math
from pathlib import Path
import statistics
import time

NETWORKS = {"LAN": {"mbps": 1000, "rtt_ms": .5}, "WAN": {"mbps": 400, "rtt_ms": 4}}


def read(path):
    return json.loads(Path(path).read_text())


def distribution(values):
    if any(value is None for value in values):
        return None
    values = sorted(values)
    if not values:
        return None
    index = .95 * (len(values) - 1)
    lo, hi = math.floor(index), math.ceil(index)
    return {"n": len(values), "median": statistics.median(values),
            "p95": values[lo] + (values[hi] - values[lo]) * (index - lo),
            "min": values[0], "max": values[-1]}


def query_phases(party):
    result = {}
    for query in party["queries"]:
        for phase, value in query["phases"].items():
            if (query["query"], phase) in result:
                raise ValueError("duplicate query phase measurement")
            result[(query["query"], phase)] = value
    for phase, value in party.get("session_phases", {}).items():
        result[(-1, phase)] = value
    for phase in ("setup", "model_prepare", "teardown"):
        if isinstance(party.get(phase), dict) and "seconds" in party[phase]:
            result[(-1, phase)] = party[phase]
    return result


def round_trace_mode(record):
    mode = record.get("round_trace", "on")
    if mode not in ("on", "off"):
        raise ValueError("round_trace must explicitly be on or off")
    return mode


def validate_paired_measurements(parties, status, graph=None):
    if len(parties) != 2:
        raise ValueError("complete paired evidence is required")
    queries = status["queries"]
    if not isinstance(queries, int) or isinstance(queries, bool) or queries < 1:
        raise ValueError("positive integer query coverage is required")
    if graph is not None:
        if not graph.get("nodes"):
            raise ValueError("complete full-graph nodes are required")
        for key in ("model", "batch", "bits", "frac"):
            if key in graph and graph[key] != status[key]:
                raise ValueError("graph identity differs from recorded plan")
        expected_layers = [(query, index, node["name"], node["op"])
                           for query in range(queries) for index, node in enumerate(graph["nodes"])]
    for role, party in enumerate(parties, 1):
        if party.get("status") != "complete" or party.get("role") != role:
            raise ValueError("incomplete party status or role")
        if any(party.get(key) != status[key] for key in ("model", "backend", "bits", "frac")):
            raise ValueError("party method/model/numeric identity differs from recorded plan")
        if [query.get("query") for query in party.get("queries", [])] != list(range(queries)):
            raise ValueError("party query coverage differs from recorded plan")
        if not status.get("diagnostic", False) and party.get("method", {}).get("diagnostic") is not False:
            raise ValueError("formal measurement requires nondiagnostic protocol evidence")
        if "optimization_profile" in status and party.get("method", {}).get("optimization_profile") != status["optimization_profile"]:
            raise ValueError("party optimization profile differs from recorded plan")
        if status["backend"] == "SCI_preprocessed":
            if party.get("method", {}).get("prepared") is not True:
                raise ValueError("preprocessed SCI must confirm its preparation mode")
            for query in party["queries"]:
                if query.get("native_generation_epoch_delta") != 0:
                    raise ValueError("preprocessed SCI generated new material online")
                written, consumed = query.get("native_spool_records"), query.get("native_spool_consumed")
                if not isinstance(written, int) or written < 0 or consumed != written:
                    raise ValueError("preprocessed SCI did not consume its recorded material inventory")
        if graph is not None:
            layers = [(layer.get("query"), layer.get("layer"), layer.get("name"), layer.get("op"))
                      for layer in party.get("layers", [])]
            if layers != expected_layers:
                raise ValueError("complete ordered full-graph layer measurements are required")
    by_party = [query_phases(party) for party in parties]
    required = {(-1, "setup")} | {(query, phase) for query in range(queries)
                                   for phase in ("offline", "input", "online", "output")}
    if not required.issubset(by_party[0]) or set(by_party[0]) != set(by_party[1]):
        raise ValueError("complete matching phase metric coverage is required")
    for key in by_party[0]:
        measured = [party[key] for party in by_party]
        for value in measured:
            for metric in ("seconds", "process_cpu_seconds"):
                number = value[metric]
                if isinstance(number, bool) or not isinstance(number, (int, float)) or not math.isfinite(number) or number < 0:
                    raise ValueError("phase time must be finite and nonnegative")
            for metric in ("sent_bytes", "received_bytes"):
                number = value[metric]
                if isinstance(number, bool) or not isinstance(number, int) or number < 0:
                    raise ValueError("phase bytes must be nonnegative integers")
        if [value["sent_bytes"] for value in measured] != [value["received_bytes"] for value in measured][::-1]:
            raise ValueError("paired phase sent/received bytes do not match")
    return by_party


def measurement_row(status, directory=None):
    row = {key: status.get(key) for key in ("model", "batch", "backend", "bits", "frac", "queries",
            "diagnostic", "status", "reason", "network", "binary_sha256", "graph_sha256",
            "parameters_sha256", "input_sha256", "bridge_library_sha256", "resource_guard", "recovery_resource_guard",
            "optimization_profile", "launcher_wall_seconds", "protocol_process_wall_seconds",
            "runtime_parameters", "tile", "cpu_affinity")}
    row["round_trace"] = round_trace_mode(status)
    row["measurement_mode"] = "time_and_bytes" if row["round_trace"] == "off" else "time_bytes_and_rounds"
    if directory is not None:
        row["directory"] = str(Path(directory).resolve())
    row["samples"] = []
    return row


def measured_costs(parties, status, rounds=None, graph=None):
    by_party = validate_paired_measurements(parties, status, graph)
    row = measurement_row(status)
    row["material_storage"] = [dict(role=party["role"], queries=[{
        key: query[key] for key in ("query", "native_generation_epoch_delta", "native_spool_bytes",
            "native_spool_records", "native_spool_consumed", "beaver_spool_bytes",
            "ours_nonlinear_spool_bytes", "conversion_material_bytes") if key in query
    } for query in party["queries"]]) for party in parties]
    if rounds is not None:
        if rounds.get("status") != "passed":
            raise ValueError("causal trace must be qualified")
        trace_keys = [(value["query"], value["phase"]) for value in rounds["phases"]]
        if set(trace_keys) != set(by_party[0]) or len(trace_keys) != len(set(trace_keys)):
            raise ValueError("complete matching phase metric/trace coverage is required")
        phases = rounds["phases"]
    else:
        row["measurement_mode"] = "time_and_bytes"
        phases = [dict(query=query, phase=phase) for query, phase in by_party[0]]
    for trace in phases:
        key = trace["query"], trace["phase"]
        measured = [party[key] for party in by_party]
        sent = [value["sent_bytes"] for value in measured]
        received = [value["received_bytes"] for value in measured]
        if rounds is not None and (sent != trace["party_sent_bytes"] or received != trace["party_received_bytes"]):
            raise ValueError("phase bytes do not match causal trace")
        causal_rounds = trace["logical_rounds"] if rounds is not None else None
        if causal_rounds is not None and (isinstance(causal_rounds, bool) or not isinstance(causal_rounds, int) or causal_rounds < 0):
            raise ValueError("causal rounds must be a nonnegative integer")
        cpu = [value["process_cpu_seconds"] for value in measured]
        wall = [value["seconds"] for value in measured]
        value = {"query": key[0], "phase": key[1], "party_cpu_seconds": cpu, "party_wall_seconds": wall,
                 "cpu_seconds_max": max(cpu), "cpu_seconds_total": sum(cpu), "wall_seconds": max(wall),
                 "sent_bytes": sum(sent), "party_sent_bytes": sent, "causal_rounds": causal_rounds,
                 "networks": {}}
        for name, network in NETWORKS.items():
            transfer = 8 * value["sent_bytes"] / (network["mbps"] * 1e6) if causal_rounds is not None else None
            latency = causal_rounds * network["rtt_ms"] / 1000 if causal_rounds is not None else None
            value["networks"][name] = {"transfer_seconds": transfer, "rtt_seconds": latency,
                                       "communication_seconds": transfer + latency if causal_rounds is not None else None}
        row["samples"].append(value)
    statistical_samples = list(row["samples"])
    for query in range(status["queries"]):
        for aggregate, included in (("offline+online", ("offline", "online")),
                                    ("inference_all_phases", ("offline", "input", "online", "output"))):
            parts = [value for value in row["samples"] if value["query"] == query and value["phase"] in included]
            cpu = [sum(value["party_cpu_seconds"][party] for value in parts) for party in (0, 1)]
            wall = [sum(value["party_wall_seconds"][party] for value in parts) for party in (0, 1)]
            statistical_samples.append(dict(phase=aggregate, cpu_seconds_max=max(cpu),
                cpu_seconds_total=sum(cpu), wall_seconds=max(wall),
                sent_bytes=sum(value["sent_bytes"] for value in parts),
                causal_rounds=sum(value["causal_rounds"] for value in parts) if rounds is not None else None,
                networks={name: {"communication_seconds": sum(value["networks"][name]["communication_seconds"]
                                                            for value in parts) if rounds is not None else None} for name in NETWORKS}))
    row["statistics"] = {}
    for phase in sorted({value["phase"] for value in statistical_samples}):
        samples = [value for value in statistical_samples if value["phase"] == phase]
        stats = {key: distribution([value[key] for value in samples]) for key in
                 ("cpu_seconds_max", "cpu_seconds_total", "wall_seconds", "sent_bytes", "causal_rounds")}
        stats["networks"] = {name: distribution([value["networks"][name]["communication_seconds"]
                                                for value in samples]) for name in NETWORKS}
        row["statistics"][phase] = stats
    return row


def paired_cost_measurement(parties, status, graph):
    """Validate complete paired inference without inspecting causal trace files."""
    return measured_costs(parties, status, graph=graph)


def file_evidence(path):
    path = Path(path).resolve()
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            value.update(block)
    return dict(path=str(path), sha256=value.hexdigest())


def validate_output_audit(audit, status):
    if audit.get("status") != "passed" or audit.get("values_checked", 0) < status["queries"] * 1000:
        raise ValueError("full-network result requires a passing 1000-logit integer output audit")


def recover_cost_measurement(directory, identity, graph, oracle_directory):
    """Recover measured time/bytes into a sidecar, preserving the original run evidence."""
    directory, oracle_directory = Path(directory), Path(oracle_directory)
    plan, status = [read(directory / name) for name in ("PLAN.json", "STATUS.json")]
    for record in (plan, status):
        for key, expected in identity.items():
            actual = round_trace_mode(record) if key == "round_trace" else record.get(key)
            if actual != expected:
                raise ValueError("existing run identity differs from the requested cost recovery")
    if round_trace_mode(plan) != round_trace_mode(status):
        raise ValueError("plan/status trace identity differs")
    parties = [read(directory / ("party%d.json" % role)) for role in (1, 2)]
    row = paired_cost_measurement(parties, status, graph)
    from audit_secure import audit_output
    audit = audit_output(directory, oracle_directory)
    validate_output_audit(audit, status)
    row.update(status="complete", reason=None, original_status=status["status"],
               directory=str(directory.resolve()), output_audit=audit,
               causal_rounds_status="not_analyzed")
    artifact = dict(row, kind="paired-inference-cost-measurement", recovered_epoch=time.time(),
                    validated_graph={key: graph[key] for key in ("model", "batch", "bits", "frac", "nodes") if key in graph},
                    source_files={name: file_evidence(directory / name) for name in
                                  ("PLAN.json", "STATUS.json", "party1.json", "party2.json", "OUTPUT_AUDIT.json")},
                    oracle=str(oracle_directory.resolve()))
    (directory / "COST_MEASUREMENT.json").write_text(json.dumps(artifact, indent=2, ensure_ascii=False) + "\n")
    return row


def extract_run(directory):
    directory = Path(directory)
    status = read(directory / "STATUS.json")
    row = measurement_row(status, directory)
    sidecar = directory / "COST_MEASUREMENT.json"
    if sidecar.exists():
        artifact = read(sidecar)
        if artifact.get("kind") != "paired-inference-cost-measurement" or artifact.get("measurement_mode") != "time_and_bytes":
            raise ValueError("invalid cost measurement sidecar")
        for name in ("PLAN.json", "STATUS.json", "party1.json", "party2.json", "OUTPUT_AUDIT.json"):
            if artifact["source_files"][name] != file_evidence(directory / name):
                raise ValueError("cost measurement source changed: " + name)
        if artifact.get("round_trace") != round_trace_mode(status):
            raise ValueError("cost measurement trace identity differs from original run")
        parties = [read(directory / ("party%d.json" % role)) for role in (1, 2)]
        row.update(paired_cost_measurement(parties, status, artifact["validated_graph"]))
        audit = read(directory / "OUTPUT_AUDIT.json")
        validate_output_audit(audit, status)
        row.update(status="complete", reason=None, original_status=status["status"], output_audit=audit,
                   causal_rounds_status="not_analyzed", cost_measurement=str(sidecar.resolve()))
        return row
    if status["status"] != "complete":
        return row
    if not status["diagnostic"]:
        output_audit = read(directory / "OUTPUT_AUDIT.json")
        validate_output_audit(output_audit, status)
        row["output_audit"] = output_audit
    parties = read(directory / "PARTIES.json")
    if round_trace_mode(status) == "off":
        live_parties = [read(directory / ("party%d.json" % role)) for role in (1, 2)]
        if parties != live_parties:
            raise ValueError("paired measurement snapshot differs from current party files")
        row.update(paired_cost_measurement(parties, status, read(status["graph"])))
        row["causal_rounds_status"] = "disabled"
    else:
        row.update(measured_costs(parties, status, read(directory / "ROUNDS.json")))
    return row


def collect(root):
    runs = []
    for path in sorted((root / "performance").glob("**/STATUS.json")):
        try:
            runs.append(extract_run(path.parent))
        except (OSError, ValueError, KeyError, TypeError) as error:
            runs.append({"directory": str(path.parent), "status": "unqualified", "reason": str(error), "samples": []})
    accuracy = []
    for path in sorted((root / "accuracy").glob("**/accuracy.json")):
        value = read(path)
        identity_path = path.parent / "metrics.json"
        if identity_path.exists():
            identity = read(identity_path)
            value["identity"] = {key: identity[key] for key in
                                 ("weights", "parameters_sha256", "graph_sha256", "numeric_contract",
                                  "input_transform", "topk_tie_break")}
            for run in runs:
                if (run.get("model"), run.get("bits"), run.get("frac"), run.get("status")) == (
                        value["model"], value["bits"], value["frac"], "complete"):
                    plan = read(Path(run["directory"]) / "PLAN.json")
                    if plan["parameters_sha256"] != identity["parameters_sha256"]:
                        raise ValueError("accuracy and secure performance use different numeric parameters")
        value["source"] = str(path.resolve())
        accuracy.append(value)
    data_status = read(root / "data/MANIFEST.json") if (root / "data/MANIFEST.json").exists() else {"status": "not_prepared"}
    return {"kind": "imagenet-resnet-evidence", "network_specs": NETWORKS, "performance": runs,
            "accuracy": accuracy, "dataset": data_status,
            "network_note": "communication model only; LAN/WAN inference wall time was not measured",
            "cpu_note": "process CPU clock; user+sys accumulated over all threads; wall includes waits"}


def workbook(data, target):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    comparison = book.create_sheet("对比")
    comparison.append(["模型", "阶段", "指标", "单位", "SCI", "SCI+ours(Beaver)", "ours/SCI", "降低(%)",
                       "SCI原始目录", "ours原始目录", "基线后端"])
    qualified = [row for row in data["performance"] if row["status"] == "complete" and not row["diagnostic"]]
    identity = ("model", "batch", "bits", "frac", "binary_sha256", "graph_sha256", "parameters_sha256", "input_sha256",
                "round_trace", "measurement_mode", "optimization_profile", "runtime_parameters", "tile", "cpu_affinity")
    for baseline in (row for row in qualified if row["backend"] in ("SCI", "SCI_preprocessed")):
        for ours in (row for row in qualified if row["backend"] == "SCI_ours_beaver" and
                     all(row.get(key) == baseline.get(key) for key in identity)):
            for phase in ("setup", "offline", "input", "online", "output", "offline+online", "inference_all_phases"):
                values = [row["statistics"][phase] for row in (baseline, ours)]
                metrics = [(key, unit, [value[key]["median"] if value[key] is not None else None for value in values]) for key, unit in
                           (("cpu_seconds_max", "s"), ("cpu_seconds_total", "s"),
                            ("wall_seconds", "s"), ("sent_bytes", "B"), ("causal_rounds", "rounds"))]
                metrics += [(name + " communication projection", "s",
                             [value["networks"][name]["median"] if value["networks"][name] is not None else None
                              for value in values]) for name in NETWORKS]
                for metric, unit, (left, right) in metrics:
                    denominator = left if (phase != "offline" or baseline["backend"] == "SCI_preprocessed") and right is not None else 0
                    comparison.append([baseline["model"], phase, metric, unit, left, right,
                                       right / denominator if denominator else None,
                                       100 * (left - right) / denominator if denominator else None,
                                       baseline["directory"], ours["directory"], baseline["backend"]])
    sessions = book.create_sheet("整次运行")
    sessions.append(["模型", "后端", "batch", "环宽", "小数位", "协议进程wall(s)", "launcher wall(s)", "优化配置", "状态", "原始目录"])
    for row in data["performance"]:
        sessions.append([row.get(key) for key in ("model", "backend", "batch", "bits", "frac", "protocol_process_wall_seconds", "launcher_wall_seconds",
                                                 "optimization_profile", "status", "directory")])
    for network in NETWORKS:
        sheet = book.create_sheet(network)
        sheet.append(["模型", "方法", "环宽", "小数位", "batch", "阶段", "统计", "n", "CPU max(s)",
                      "CPU total(s)", "本机wall(s)", "两方sent(B)", "MiB", "因果轮数", "通信估算(s)", "状态", "范围",
                      "trace开关", "测量模式"])
        for row in data["performance"]:
            if row["status"] != "complete":
                sheet.append([row.get("model"), row.get("backend"), row.get("bits"), row.get("frac"), row.get("batch"),
                              None, None, None, None, None, None, None, None, None, None, row["status"], row.get("reason"),
                              row.get("round_trace"), row.get("measurement_mode")])
                continue
            for phase, stats in row["statistics"].items():
                for statistic in (("median",) if stats["wall_seconds"]["n"] == 1 else ("median", "p95")):
                    label = "single" if stats["wall_seconds"]["n"] == 1 else statistic
                    sheet.append([row["model"], row["backend"], row["bits"], row["frac"], row["batch"], phase, label,
                        stats["wall_seconds"]["n"], *[stats[key][statistic] for key in
                        ("cpu_seconds_max", "cpu_seconds_total", "wall_seconds", "sent_bytes")],
                        stats["sent_bytes"][statistic] / 2**20,
                        stats["causal_rounds"][statistic] if stats["causal_rounds"] is not None else None,
                        stats["networks"][network][statistic] if stats["networks"][network] is not None else None, row["status"],
                        "diagnostic" if row["diagnostic"] else "full network", row["round_trace"], row["measurement_mode"]])
    accuracy = book.create_sheet("精度")
    accuracy.append(["模型", "环宽", "小数位", "实际样本", "期望样本", "范围", "FP32 Top1(%)", "FP32 Top5(%)",
                     "定点Top1(%)", "定点Top5(%)", "Top1差值(pp)", "Top5差值(pp)",
                     "Top1预测一致(%)", "logit MAE", "logit RMSE", "logit最大绝对误差", "状态", "原始记录"])
    for row in data["accuracy"]:
        fp32, fixed = row.get("fp32", {}), row.get("fixed", {})
        delta = [fixed.get(key) - fp32.get(key) if fixed.get(key) is not None and fp32.get(key) is not None
                 else None for key in ("top1_pct", "top5_pct")]
        accuracy.append([row.get("model"), row.get("bits"), row.get("frac"), row.get("samples"),
            row.get("expected_samples", 50000), row.get("evidence_scope"), fp32.get("top1_pct"),
            fp32.get("top5_pct"), fixed.get("top1_pct"), fixed.get("top5_pct"), *delta,
            row.get("top1_agreement_pct"), row.get("logit_mae"), row.get("logit_rmse"),
            row.get("logit_max_abs_error"), row.get("status"), row["source"]])
    raw = book.create_sheet("逐query")
    raw.append(["模型", "方法", "query", "阶段", "wall方1", "wall方2", "CPU方1", "CPU方2", "sent方1", "sent方2",
                "因果轮数", "LAN通信估算(s)", "WAN通信估算(s)", "原始目录", "trace开关", "测量模式"])
    for row in data["performance"]:
        for value in row["samples"]:
            raw.append([row["model"], row["backend"], value["query"], value["phase"], *value["party_wall_seconds"],
                *value["party_cpu_seconds"], *value["party_sent_bytes"], value["causal_rounds"],
                *[value["networks"][network]["communication_seconds"] for network in NETWORKS], row["directory"],
                row["round_trace"], row["measurement_mode"]])
    notes = book.create_sheet("说明")
    notes.append(["项目", "说明"])
    for key, value in {
        "计算时间": data["cpu_note"], "网络": data["network_note"],
        "计时范围": "CPU和本机wall包含协议与材料文件读写；trace=on另含轨迹记录开销。trace=off仅测耗时和字节，轮数与网络投影留空。部分旧计时与GPU精度回放重叠，CPU亲和性不重叠，未固定频率",
        "内存保护": "恢复后新运行使用进程树cgroup限制2 GiB RAM加512 MiB swap；resource_guard为协议限制，recovery_resource_guard仅约束事后分析。限制前后CPU/wall需结合swap与回收情况解释",
        "LAN": "1000 Mbps, RTT 0.5 ms", "WAN": "400 Mbps, RTT 4 ms",
        "通信模型": "8 * 两方sent / bandwidth + causal rounds * RTT；不加入本机wall",
        "数值": "Conv-BN折叠；环累加，signed floor截断，之后加bias；ReLU精确；AvgPool环求和后floor除法",
        "方法": "SCI原生OT在线生成材料；SCI_preprocessed用相同原生OT/triple提前生成并写私有磁盘；SCI_ours_beaver使用优化FSS与矩阵Beaver。AvgPool共用精确公开除法分解",
        "精度": "正式ImageNet需50000张，原始1000类；diagnostic或partial不能当正式精度",
        "未完成": "空白表示没有完成或不适用，不能当零开销或零准确率",
        "样本数": "single表示该配置仅完整执行1次；不据此报告延迟方差或稳定性",
        "计数": "双方sent之和，一次消息不重复加received；轮数来自完整阶段因果轨迹",
        "合计": "offline+online只含两个阶段；inference_all_phases另含input/output，均逐方CPU/wall求和后取max，不含setup。协议进程wall包括setup、图加载及结果保存；launcher wall另含进程退出后的校验及trace分析",
        "比值": "仅配对同图、同输入、同参数、同binary、同trace开关与同测量模式的完整SCI与Beaver配置；降低为负表示开销增加",
        "离线比值": "非预处理SCI离线仅有控制调用，其离线比值留空；SCI_preprocessed按实测离线材料生成成本比较",
    }.items():
        notes.append([key, value])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="34495E")
        for index, cells in enumerate(sheet.columns, 1):
            width = max(len(str(cell.value or "")) for cell in cells) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(60, max(12, width))
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
    book.save(target)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, required=True)
    args = parser.parse_args()
    root = args.directory.resolve()
    data = collect(root)
    (root / "METRICS.json").write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    workbook(data, root / "REPORT.xlsx")
    print(json.dumps({"performance_rows": len(data["performance"]), "accuracy_rows": len(data["accuracy"]),
                      "excel": str(root / "REPORT.xlsx")}))


if __name__ == "__main__":
    main()
