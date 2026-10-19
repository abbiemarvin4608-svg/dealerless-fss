"""Check measured time/byte evidence without manufacturing communication rounds."""
import copy
import json
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch

import summarize


class CostMeasurementTests(unittest.TestCase):
    def test_runtime_parameters_do_not_inherit_protocol_overrides(self):
        from run_secure import RUNTIME_PARAMETERS, runtime_environment
        parent = {"SEGFSS_PREFETCH_BYTES": "0", "SEGFSS_UNKNOWN_FLAG": "1", "SDCF_ROUND_TRACE": "old", "PATH": "/bin"}
        environment = runtime_environment(parent)
        self.assertEqual(environment["SEGFSS_PREFETCH_BYTES"], RUNTIME_PARAMETERS["SEGFSS_PREFETCH_BYTES"])
        self.assertNotIn("SEGFSS_UNKNOWN_FLAG", environment)
        self.assertNotIn("SDCF_ROUND_TRACE", environment)
        self.assertEqual(environment["PATH"], "/bin")

    def test_child_environment_explicitly_disables_inherited_trace(self):
        from run_secure import child_environment
        parent = {"SDCF_ROUND_TRACE": "/previous/run.trace.jsonl", "OMP_NUM_THREADS": "1"}
        original = dict(parent)
        output = Path("/current/run")
        disabled = child_environment(parent, output, 1, "off")
        self.assertNotIn("SDCF_ROUND_TRACE", disabled)
        self.assertEqual(disabled["OMP_NUM_THREADS"], "1")
        enabled = child_environment(parent, output, 2, "on")
        self.assertEqual(enabled["SDCF_ROUND_TRACE"], str(output / "party2.trace.jsonl"))
        self.assertEqual(enabled["OMP_NUM_THREADS"], "1")
        self.assertEqual(parent, original)

    def write(self, path, value):
        path.write_text(json.dumps(value))

    def fixture(self, directory, *, diagnostic=True, round_trace="off"):
        graph = dict(nodes=[dict(name="fc", op="linear")])
        graph_path = directory / "graph.json"
        self.write(graph_path, graph)
        status = dict(status="complete", model="cost_fixture", batch=1, backend="SCI",
                      bits=32, frac=12, queries=1, diagnostic=diagnostic,
                      reason=None, graph=str(graph_path))
        if round_trace is not None:
            status["round_trace"] = round_trace
        parties = []
        for index in range(2):
            zero = dict(seconds=0., process_cpu_seconds=0., sent_bytes=0, received_bytes=0)
            online = dict(seconds=[.3, .4][index], process_cpu_seconds=[.8, .7][index],
                          sent_bytes=[120, 80][index], received_bytes=[80, 120][index])
            offline = dict(seconds=[.05, .1][index], process_cpu_seconds=[.1, .5][index],
                           sent_bytes=[20, 30][index], received_bytes=[30, 20][index])
            parties.append(dict(status="complete", role=index + 1, model=status["model"],
                backend=status["backend"], bits=32, frac=12, method=dict(diagnostic=diagnostic),
                setup=zero, queries=[dict(query=0, phases=dict(online=online, offline=offline,
                                                             input=zero, output=zero),
                                         output_codes=[0] * 1000)],
                layers=[dict(query=0, layer=0, name="fc", op="linear")]))
        self.write(directory / "STATUS.json", status)
        self.write(directory / "PARTIES.json", parties)
        for role, party in enumerate(parties, 1):
            self.write(directory / ("party%d.json" % role), party)
        if not diagnostic:
            self.write(directory / "OUTPUT_AUDIT.json", dict(status="passed", values_checked=1000))
        return status, parties, graph

    def test_explicit_trace_off_preserves_measured_costs_and_null_rounds(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory)
            row = summarize.extract_run(directory)
            self.assertEqual(row["status"], "complete")
            self.assertEqual(row["round_trace"], "off")
            online = next(value for value in row["samples"] if value["phase"] == "online")
            self.assertEqual(online["wall_seconds"], .4)
            self.assertEqual(online["cpu_seconds_max"], .8)
            self.assertEqual(online["cpu_seconds_total"], 1.5)
            self.assertEqual(online["sent_bytes"], 200)
            for sample in row["samples"]:
                self.assertIsNone(sample["causal_rounds"])
                for network in summarize.NETWORKS:
                    for metric in ("transfer_seconds", "rtt_seconds", "communication_seconds"):
                        self.assertIsNone(sample["networks"][network][metric])
            total = row["statistics"]["offline+online"]
            self.assertAlmostEqual(total["cpu_seconds_max"]["median"], 1.2)
            self.assertAlmostEqual(total["cpu_seconds_total"]["median"], 2.1)
            self.assertAlmostEqual(total["wall_seconds"]["median"], .5)
            self.assertEqual(total["sent_bytes"]["median"], 250)
            self.assertIsNone(total["causal_rounds"])
            self.assertTrue(all(value is None for value in total["networks"].values()))
            self.assertAlmostEqual(row["statistics"]["inference_all_phases"]["wall_seconds"]["median"], .5)
            self.assertFalse((directory / "ROUNDS.json").exists())

    def test_trace_off_rejects_incomplete_phases_bytes_and_layers(self):
        for defect in ("phase", "bytes", "layer"):
            with self.subTest(defect=defect), tempfile.TemporaryDirectory() as temporary:
                directory = Path(temporary)
                status, parties, graph = self.fixture(directory)
                if defect == "phase":
                    del parties[1]["queries"][0]["phases"]["input"]
                elif defect == "bytes":
                    parties[0]["queries"][0]["phases"]["online"]["sent_bytes"] += 1
                else:
                    parties[1]["layers"] = []
                self.write(directory / "PARTIES.json", parties)
                with self.assertRaises(ValueError):
                    summarize.paired_cost_measurement(parties, status, graph)
                with self.assertRaises(ValueError):
                    summarize.extract_run(directory)

    def test_missing_trace_metadata_never_silently_downgrades_to_costs(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory, round_trace=None)
            with self.assertRaises((OSError, ValueError)):
                summarize.extract_run(directory)

    def test_preprocessed_sci_requires_no_online_generation_and_full_consumption(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            status, parties, graph = self.fixture(directory)
            status.update(backend="SCI_preprocessed", optimization_profile="optimized")
            for party in parties:
                party["backend"] = "SCI_preprocessed"
                party["method"].update(prepared=True, optimization_profile="optimized")
                party["queries"][0].update(native_generation_epoch_delta=0,
                    native_spool_records=12, native_spool_consumed=12)
            summarize.paired_cost_measurement(parties, status, graph)
            for key, value in (("native_generation_epoch_delta", 1), ("native_spool_consumed", 11)):
                changed = copy.deepcopy(parties)
                changed[0]["queries"][0][key] = value
                with self.assertRaises(ValueError):
                    summarize.paired_cost_measurement(changed, status, graph)

    def test_trace_off_rejects_stale_party_snapshot(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            _, parties, _ = self.fixture(directory, diagnostic=False)
            parties[0]["queries"][0]["phases"]["online"]["sent_bytes"] += 1
            parties[1]["queries"][0]["phases"]["online"]["received_bytes"] += 1
            for role, party in enumerate(parties, 1):
                self.write(directory / ("party%d.json" % role), party)
            with self.assertRaises(ValueError):
                summarize.extract_run(directory)

    def test_workbook_keeps_rounds_and_network_projections_blank(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory, diagnostic=False)
            baseline = summarize.extract_run(directory)
            ours = copy.deepcopy(baseline)
            ours.update(backend="SCI_ours_beaver", directory=str(directory / "ours"))
            target = directory / "costs.xlsx"
            summarize.workbook(dict(performance=[baseline, ours], accuracy=[],
                                    cpu_note="CPU", network_note="No round projection"), target)
            book = load_workbook(target, data_only=True)
            try:
                for name in summarize.NETWORKS:
                    rows = list(book[name].iter_rows(min_row=2, values_only=True))
                    self.assertTrue(rows)
                    for row in rows:
                        self.assertIsNone(row[13])
                        self.assertIsNone(row[14])
                        self.assertEqual(row[15], "complete")
                for row in book["逐query"].iter_rows(min_row=2, values_only=True):
                    self.assertEqual(row[10:13], (None, None, None))
                comparison = list(book["对比"].iter_rows(min_row=2, values_only=True))
                self.assertTrue(any(row[2] == "sent_bytes" and row[4] is not None for row in comparison))
                for row in comparison:
                    if row[2] == "causal_rounds" or "communication projection" in row[2]:
                        self.assertEqual(row[4:8], (None, None, None, None))
            finally:
                book.close()

    def test_workbook_never_compares_trace_on_with_trace_off_costs(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory, diagnostic=False)
            baseline = summarize.extract_run(directory)
            ours = copy.deepcopy(baseline)
            ours.update(backend="SCI_ours_beaver", round_trace="on",
                        directory=str(directory / "ours"))
            target = directory / "mixed.xlsx"
            summarize.workbook(dict(performance=[baseline, ours], accuracy=[],
                                    cpu_note="CPU", network_note="No round projection"), target)
            book = load_workbook(target, data_only=True)
            try:
                self.assertEqual(book["对比"].max_row, 1)
            finally:
                book.close()

    def recovery_fixture(self, root):
        directory = root / "run"
        directory.mkdir()
        status, parties, graph = self.fixture(directory, diagnostic=False, round_trace="on")
        status.update(status="running", binary_sha256="fixture-binary")
        self.write(directory / "STATUS.json", status)
        self.write(directory / "PLAN.json", status)
        for role, party in enumerate(parties, 1):
            self.write(directory / ("party%d.json" % role), party)
            (directory / ("party%d.trace.jsonl" % role)).write_text("original protocol trace\n")
        oracle = root / "oracle"
        oracle.mkdir()
        self.write(oracle / "index.json", dict(model=status["model"], bits=32, frac=12,
            tensors=dict(fc=dict(file="fc.bin", elements=1000, shape=[1, 1000]))))
        (oracle / "fc.bin").write_bytes(struct.pack("<1000Q", *([0] * 1000)))
        identity = {key: status[key] for key in ("model", "batch", "backend", "bits", "frac",
                                               "queries", "diagnostic", "binary_sha256")}
        return directory, identity, graph, oracle

    def test_recover_costs_validates_outputs_and_preserves_protocol_evidence(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory, identity, graph, oracle = self.recovery_fixture(Path(temporary))
            originals = {name: (directory / name).read_bytes() for name in
                         ("STATUS.json", "PLAN.json", "party1.json", "party2.json",
                          "party1.trace.jsonl", "party2.trace.jsonl")}
            with patch("round_trace_stream.analyze_pair", side_effect=AssertionError("trace analysis must not run")):
                row = summarize.recover_cost_measurement(directory, identity, graph, oracle)
                extracted = summarize.extract_run(directory)
            self.assertEqual(row["status"], "complete")
            self.assertEqual(row["original_status"], "running")
            self.assertEqual(row["round_trace"], "on")
            self.assertEqual(row["measurement_mode"], "time_and_bytes")
            self.assertEqual(row["samples"], extracted["samples"])
            self.assertEqual(row["statistics"], extracted["statistics"])
            self.assertTrue((directory / "COST_MEASUREMENT.json").is_file())
            self.assertFalse((directory / "ROUNDS.json").exists())
            for name, original in originals.items():
                self.assertEqual((directory / name).read_bytes(), original, name)

    def test_wrong_logits_cannot_be_recovered_as_complete_costs(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory, identity, graph, oracle = self.recovery_fixture(Path(temporary))
            party_path = directory / "party2.json"
            party = json.loads(party_path.read_text())
            party["queries"][0]["output_codes"][17] = 1
            self.write(party_path, party)
            original = (directory / "STATUS.json").read_bytes()
            with self.assertRaises(ValueError):
                summarize.recover_cost_measurement(directory, identity, graph, oracle)
            self.assertEqual((directory / "STATUS.json").read_bytes(), original)
            self.assertFalse((directory / "COST_MEASUREMENT.json").exists())


if __name__ == "__main__":
    unittest.main()
