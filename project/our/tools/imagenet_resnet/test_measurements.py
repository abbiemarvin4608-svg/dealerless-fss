"""Check that result aggregation rejects invalid communication evidence."""
import json
from pathlib import Path
import tempfile
import unittest

from summarize import extract_run, workbook


class MeasurementTests(unittest.TestCase):
    def fixture(self, directory):
        status = dict(status="complete", model="measurement_fixture", batch=1, backend="SCI",
                      bits=32, frac=12, queries=1, diagnostic=True, reason=None)
        parties = []
        for index in range(2):
            phase = dict(seconds=[.3, .4][index], process_cpu_seconds=[.8, .7][index],
                         sent_bytes=[120, 80][index], received_bytes=[80, 120][index])
            zero = dict(seconds=0., process_cpu_seconds=0., sent_bytes=0, received_bytes=0)
            parties.append(dict(status="complete", role=index + 1, model=status["model"], backend=status["backend"],
                bits=status["bits"], frac=status["frac"], setup=zero,
                queries=[dict(query=0, phases={"online": phase, "offline": zero, "input": zero, "output": zero})]))
        phases = [dict(query=0, phase="online", logical_rounds=7, party_sent_bytes=[120, 80], party_received_bytes=[80, 120])]
        phases += [dict(query=0, phase=name, logical_rounds=0, party_sent_bytes=[0, 0], party_received_bytes=[0, 0])
                   for name in ("offline", "input", "output")]
        phases += [dict(query=-1, phase="setup", logical_rounds=0, party_sent_bytes=[0, 0], party_received_bytes=[0, 0])]
        rounds = dict(status="passed", phases=phases)
        for name, value in (("STATUS.json", status), ("PARTIES.json", parties), ("ROUNDS.json", rounds)):
            (directory / name).write_text(json.dumps(value))
        return parties

    def test_compute_bytes_and_rounds_remain_separate(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory)
            value = extract_run(directory)["samples"][0]
            self.assertEqual(value["wall_seconds"], .4)
            self.assertEqual(value["cpu_seconds_max"], .8)
            self.assertEqual(value["cpu_seconds_total"], 1.5)
            self.assertEqual(value["sent_bytes"], 200)
            self.assertEqual(value["causal_rounds"], 7)
            self.assertAlmostEqual(value["networks"]["WAN"]["communication_seconds"], 200 * 8 / 400e6 + 7 * .004)

    def test_trace_bytes_mismatch_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            parties = self.fixture(directory)
            parties[0]["queries"][0]["phases"]["online"]["sent_bytes"] += 1
            (directory / "PARTIES.json").write_text(json.dumps(parties))
            with self.assertRaisesRegex(ValueError, "bytes"):
                extract_run(directory)

    def test_phase_sum_takes_max_after_adding_each_party(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            parties = self.fixture(directory)
            for index, party in enumerate(parties):
                party["queries"][0]["phases"]["offline"] = dict(
                    seconds=[.05, .1][index], process_cpu_seconds=[.1, .5][index],
                    sent_bytes=[20, 30][index], received_bytes=[30, 20][index])
            (directory / "PARTIES.json").write_text(json.dumps(parties))
            rounds = json.loads((directory / "ROUNDS.json").read_text())
            for phase in rounds["phases"]:
                if phase["phase"] == "offline":
                    phase.update(logical_rounds=3, party_sent_bytes=[20, 30], party_received_bytes=[30, 20])
            (directory / "ROUNDS.json").write_text(json.dumps(rounds))
            total = extract_run(directory)["statistics"]["offline+online"]
            self.assertAlmostEqual(total["cpu_seconds_max"]["median"], 1.2)
            self.assertAlmostEqual(total["cpu_seconds_total"]["median"], 2.1)
            self.assertAlmostEqual(total["wall_seconds"]["median"], .5)
            self.assertEqual(total["sent_bytes"]["median"], 250)
            self.assertEqual(total["causal_rounds"]["median"], 10)

    def test_missing_phase_and_query_are_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory)
            rounds = json.loads((directory / "ROUNDS.json").read_text())
            rounds["phases"] = rounds["phases"][:-1]
            (directory / "ROUNDS.json").write_text(json.dumps(rounds))
            with self.assertRaisesRegex(ValueError, "coverage"):
                extract_run(directory)
            parties = self.fixture(directory)
            parties[1]["queries"] = []
            (directory / "PARTIES.json").write_text(json.dumps(parties))
            with self.assertRaisesRegex(ValueError, "query"):
                extract_run(directory)

    def test_timeout_is_missing_in_excel(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            self.fixture(directory)
            status = json.loads((directory / "STATUS.json").read_text())
            status.update(status="timeout", reason="budget")
            (directory / "STATUS.json").write_text(json.dumps(status))
            row = extract_run(directory)
            self.assertEqual(row["samples"], [])
            workbook(dict(performance=[row], accuracy=[], cpu_note="CPU", network_note="model only"), directory / "test.xlsx")
            book = load_workbook(directory / "test.xlsx", data_only=True)
            self.assertIsNone(book["LAN"].cell(2, 9).value)
            self.assertEqual(book["LAN"].cell(2, 16).value, "timeout")


if __name__ == "__main__":
    unittest.main()
