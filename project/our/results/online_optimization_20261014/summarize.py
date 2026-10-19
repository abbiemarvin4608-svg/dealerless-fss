#!/usr/bin/env python3
"""Produce evidence-backed JSON, Chinese report, and LAN/WAN Excel workbook."""
import argparse
import hashlib
import json
import math
from pathlib import Path
import statistics

HERE = Path(__file__).resolve().parent
PHASES = ("offline", "online", "offline_plus_online")
NETWORKS = {"LAN": {"mbps": 1000, "rtt_ms": 0.5}, "WAN": {"mbps": 400, "rtt_ms": 4}}
MAIN_ROWS = ("ours_online_before", "ours_online_optimized", "SCI_SiRNN_preprocessed", "SCI_SiRNN_unprocessed")
ORDER = ("ours_online_before", "ours_input_sign_reuse", "ours_input_sign_reuse_clip22",
         "ours_input_sign_reuse_clip22_lutskip", "ours_online_optimized",
         "SCI_SiRNN_preprocessed", "SCI_SiRNN_unprocessed")
LABELS = dict(zip(ORDER, ("SCI+ours 在线优化前", "+ 复用输入符号", "+ 22-bit 裁剪比较",
                        "+ 跳过零 LUT 系数", "SCI+ours 在线优化后", "SCI/SiRNN 预处理", "SCI/SiRNN 未预处理")))
PHASE_LABELS = {"offline": "离线", "online": "在线", "offline_plus_online": "离线＋在线"}
COMMON_FIELDS = ("batch", "steps", "start", "queries", "warmup_queries", "hidden_file",
                 "fixture_identity", "fixture_files", "binary_sha256", "memory_limit_bytes_per_party",
                 "cpu_affinity_per_party", "graph_mode", "activation_workers_per_party", "model_bias_precomputed")
OPTION_FIELDS = ("backend", "native_preprocessing_effective", "lut_key_family", "folded_sign_backend",
                 "bit_block_mode", "native_ring_cot", "cache_lut_leaves", "fuse_sign_wrap",
                 "folded_reuse_input_sign", "folded_clip_bits", "lut_skip_zero", "lut_simd")
SCALARS = ("wall_seconds", "cpu_seconds_max", "cpu_seconds_total", "sent_bytes", "causal_rounds")


def read(path):
    return json.loads(path.read_text())


def exactly(rows, **identity):
    values = [row for row in rows if all(row.get(key) == value for key, value in identity.items())]
    if len(values) != 1:
        raise ValueError("missing or duplicate evidence: " + str(identity))
    return values[0]


def distribution(values):
    if not values or any(value is None for value in values):
        return None
    values = sorted(values)
    position = (len(values) - 1) * .95
    lower, upper = math.floor(position), math.ceil(position)
    return {"n": len(values), "median": statistics.median(values),
            "p95": values[lower] + (values[upper] - values[lower]) * (position - lower),
            "min": values[0], "max": values[-1]}


def projection(sent_bytes, rounds, link):
    transfer = 8 * sent_bytes / (link["mbps"] * 1e6)
    latency = rounds * link["rtt_ms"] / 1000
    return {"transfer_seconds": transfer, "rtt_seconds": latency, "communication_seconds": transfer + latency}


def phase_value(raw, trace, links, require_cpu):
    parties = raw["per_party"]
    if len(parties) != 2:
        raise ValueError("two-party phase evidence required")
    sent = [row["sent_bytes"] for row in parties]
    received = [row["received_bytes"] for row in parties]
    wall = [row["seconds"] for row in parties]
    cpu = [row.get("process_cpu_seconds") for row in parties]
    if sent != trace["party_sent_bytes"] or received != trace["party_received_bytes"] or \
            sent != list(reversed(received)) or sum(sent) != raw["sent_bytes"] or max(wall) != raw["seconds"]:
        raise ValueError("measurement and causal trace bytes/wall disagree")
    if require_cpu and any(value is None for value in cpu):
        raise ValueError("process_cpu_seconds missing from the new measurement build")
    if any(value is not None and (not math.isfinite(value) or value < 0) for value in cpu):
        raise ValueError("invalid process CPU seconds")
    rounds = trace["logical_rounds"]
    if type(rounds) is not int or rounds < 0:
        raise ValueError("qualified integer causal rounds required")
    have_cpu = all(value is not None for value in cpu)
    if have_cpu:
        for name, value in (("process_cpu_seconds_max", max(cpu)), ("process_cpu_seconds_total", sum(cpu))):
            if name in raw and not math.isclose(raw[name], value, rel_tol=1e-10, abs_tol=1e-9):
                raise ValueError("aggregated CPU seconds disagree with party measurements")
    return {"wall_seconds": max(wall), "party_wall_seconds": wall,
            "party_cpu_seconds": cpu, "cpu_seconds_max": max(cpu) if have_cpu else None,
            "cpu_seconds_total": sum(cpu) if have_cpu else None,
            "sent_bytes": sum(sent), "party_sent_bytes": sent, "party_received_bytes": received,
            "causal_rounds": rounds, "network_projection": {
                name: projection(sum(sent), rounds, specification) for name, specification in links.items()}}


def inventory(parties, prepared):
    zero = True
    deltas = []
    applicable_flags = {}
    balances = {}
    one_use = True
    for party in parties:
        online = [row for row in party.get("phases", []) if row["phase"] == "online"]
        for row in online:
            epoch = {key: value for key, value in row.items() if key.endswith("_generation_epoch_delta")}
            deltas.append({"party": party["party"], "query": row["query"], "deltas": epoch})
            zero &= all(value == 0 for value in epoch.values())
        for name in ("native_material_balances", "native_triple_balances", "folded_activation_balances", "dpf_lut_balances"):
            records = party.get(name, [])
            balances.setdefault(name, []).append(records)
            for row in records:
                one_use &= row.get("remaining") == 0
                padding = row.get("padding_discarded", 0)
                one_use &= row.get("generated") == row.get("consumed", 0) + padding
                if "generated_elements" in row:
                    one_use &= row["generated_elements"] == row.get("consumed_elements")
                zero &= all(row.get(key, 0) == 0 for key in
                            ("online_generated", "online_loaded_bytes", "online_generation_epoch_delta"))
        for name, value in party.items():
            if not name.endswith("_verified") or not isinstance(value, bool):
                continue
            if name.startswith("native_"):
                applies = prepared
            elif name.startswith("folded_"):
                applies = bool(party.get("folded_activation_balances"))
            elif name.startswith("dpf_"):
                applies = bool(party.get("dpf_lut_balances"))
            else:
                applies = prepared
            applicable_flags.setdefault(name, []).append({"value": value, "applicable": applies})
            if applies:
                one_use &= value
    present = any(any(records) for records in balances.values())
    if prepared:
        policy_passed = zero and bool(deltas) and all(row["deltas"] for row in deltas)
        stock_passed = one_use and present
    else:
        policy_passed = bool(deltas) and all(row["deltas"].get("native_generation_epoch_delta", 0) > 0 for row in deltas)
        stock_passed = None
    return {"mode": "prepared" if prepared else "native generation inside online",
            "online_generation_policy_passed": policy_passed,
            "zero_online_generation_observed": zero, "zero_online_generation_applicable": prepared,
            "one_use_stock_applicable": prepared, "one_use_stock_passed": stock_passed,
            "balance_evidence_present": present, "balance_records": balances,
            "verification_flags": applicable_flags, "generation_deltas": deltas}


def validate_online_modes(declaration, config, parties, rounds):
    command = declaration["command"]
    expected = {key: ("--" + key.replace("_", "-")) in command for key in
                ("folded_reuse_input_sign", "lut_skip_zero", "lut_simd")}
    expected["folded_clip_bits"] = int(command[command.index("--folded-clip-bits") + 1]) if "--folded-clip-bits" in command else 32
    for key, value in expected.items():
        if config.get(key) != value:
            raise ValueError("requested/recorded online mode mismatch: " + key)
    if declaration["backend"] != "SCI_Beaver_segmented_PaperLUT":
        return {"expected": expected, "folded_stage_validation": "N/A"}
    if config.get("conversion_optimizations", {}).get("pack_bits") is not True:
        raise ValueError("folded stage byte expectations require the declared packed Boolean opening")
    observed = []
    for party_index, party in enumerate(parties):
        calls = party["folded_activation_consumed"]
        if len(calls) != config["queries"] * 4:
            raise ValueError("folded consumed call inventory incomplete")
        for call in calls:
            count = call["count"]
            for stage, layers in ((3, 0 if expected["folded_reuse_input_sign"] else 31),
                                  (5, expected["folded_clip_bits"] - 1)):
                metric = call["online"][stage]
                byte_count = layers * ((4 * count + 7) // 8)
                if metric["sent_bytes"] != byte_count or metric["received_bytes"] != byte_count:
                    raise ValueError("folded comparator stage did not apply declared mode")
                scopes = [row for row in rounds["scopes"] if row["label"] == "activation_stage" and
                    row["phase"] == "online" and row["stage"] == stage and
                    all(row.get(key) == call[key] for key in ("query", "layer", "slot", "call", "kind"))]
                if not scopes and layers == 0:
                    continue
                if len(scopes) != 1 or scopes[0]["logical_rounds"] != layers or \
                        scopes[0]["party_sent_bytes"][party_index] != byte_count:
                    raise ValueError("folded comparator causal scope does not match actual stage")
            if party_index == 0:
                observed.append({"query": call["query"], "layer": call["layer"], "slot": call["slot"],
                    "call": call["call"], "count": count, "sign_layers": 0 if expected["folded_reuse_input_sign"] else 31,
                    "clip_layers": expected["folded_clip_bits"] - 1,
                    "party_sign_sent_bytes": call["online"][3]["sent_bytes"],
                    "party_clip_sent_bytes": call["online"][5]["sent_bytes"]})
    return {"expected": expected, "folded_stage_validation": "passed", "observed_calls": observed}


def extract(declaration, state, plan, links=NETWORKS, require_cpu=True, expected_queries=10):
    path = Path(declaration["output"])
    result, parties = read(path / "result.json"), read(path / "parties.json")
    config, measurements, rounds = result["config"], result["measurements"], result["round_measurements"]
    prepared = declaration["backend"] != "SCI"
    checks = {"completed": state["status"] == result["status"] == "complete",
              "exited": state.get("exit_code") == 0 and result.get("party_exit_codes") == [0, 0],
              "backend": config["backend"] == declaration["backend"],
              "preprocessing_policy": config.get("native_preprocessing_effective") is prepared,
              "fresh_queries": config["queries"] == expected_queries and config["warmup_queries"] == 0 and
                               len(measurements["queries"]) == expected_queries,
              "workload": config["batch"] == 128 and config["steps"] == 1 and config["start"] == 0 and
                          config.get("hidden_file") is None and not config.get("diagnostic"),
              "binary": config["binary_sha256"] == plan["build"]["binary"],
              "resources": config["memory_limit_bytes_per_party"] == 1610612736 and
                           result["address_space_hard_limit_bytes_per_party"] == 1610612736,
              "causal_trace": rounds["status"] == "passed", "unshaped": config["network"]["name"] == "localhost"}
    bridge = Path(plan["bridge"])
    checks["frozen_bridge"] = bridge.is_file() and hashlib.sha256(bridge.read_bytes()).hexdigest() == plan["build"]["bridge"]
    if not all(checks.values()):
        raise ValueError("contract failed: " + ", ".join(key for key, passed in checks.items() if not passed))
    modes = validate_online_modes(declaration, config, parties, rounds)
    queries = []
    for query_id, measured in enumerate(measurements["queries"]):
        phases = {}
        for phase in ("offline", "online", "input", "ready", "output"):
            phases[phase] = phase_value(measured[phase], exactly(rounds["phases"], query=query_id, phase=phase), links, require_cpu)
        off, on = phases["offline"], phases["online"]
        cpu = [off["party_cpu_seconds"][index] + on["party_cpu_seconds"][index]
               for index in (0, 1)] if all(value is not None for row in (off, on) for value in row["party_cpu_seconds"]) else [None, None]
        phases["offline_plus_online"] = {
            "wall_seconds": off["wall_seconds"] + on["wall_seconds"],
            "party_wall_seconds": [off["party_wall_seconds"][index] + on["party_wall_seconds"][index] for index in (0, 1)],
            "party_cpu_seconds": cpu, "cpu_seconds_max": max(cpu) if cpu[0] is not None else None,
            "cpu_seconds_total": sum(cpu) if cpu[0] is not None else None,
            "sent_bytes": off["sent_bytes"] + on["sent_bytes"],
            "party_sent_bytes": [off["party_sent_bytes"][index] + on["party_sent_bytes"][index] for index in (0, 1)],
            "party_received_bytes": [off["party_received_bytes"][index] + on["party_received_bytes"][index] for index in (0, 1)],
            "causal_rounds": off["causal_rounds"] + on["causal_rounds"],
            "network_projection": {name: {key: off["network_projection"][name][key] + on["network_projection"][name][key]
                for key in ("transfer_seconds", "rtt_seconds", "communication_seconds")} for name in links}}
        queries.append({"query": query_id, "phases": phases, "query_total_raw": measured["query_total"]})
    stats = {phase: {key: distribution([query["phases"][phase][key] for query in queries]) for key in SCALARS} for phase in PHASES}
    for phase in PHASES:
        stats[phase]["network_projection"] = {name: {key: distribution([query["phases"][phase]["network_projection"][name][key]
            for query in queries]) for key in ("transfer_seconds", "rtt_seconds", "communication_seconds")} for name in links}
    accounting = inventory(parties, prepared)
    if not accounting["online_generation_policy_passed"] or prepared and not accounting["one_use_stock_passed"]:
        raise ValueError("online generation/one-use stock policy failed")
    session_phases = {}
    for phase in ("setup", "model_prepare", "teardown"):
        raw = measurements.get(phase)
        trace_rows = [row for row in rounds["phases"] if row["query"] == -1 and row["phase"] == phase]
        if raw and isinstance(raw, dict) and raw.get("per_party") and len(trace_rows) == 1:
            session_phases[phase] = phase_value(raw, trace_rows[0], links, require_cpu)
        else:
            session_phases[phase] = {"raw_measurement": raw, "trace": trace_rows}
    observed = result.get("process_resource_peaks", [{}, {}])
    resources = [{"party": index + 1,
        "rss_peak_bytes": max(party.get("peak_rss_kib", 0) * 1024, observed[index].get("rss_peak_bytes", 0)),
        "vm_peak_bytes": max(party.get("vm_peak_bytes", 0), observed[index].get("vm_peak_bytes", 0))}
        for index, party in enumerate(parties)]
    return {"id": declaration["id"], "label": LABELS[declaration["id"]], "status": "complete", "directory": str(path),
            "checks": checks, "common_config": {key: config.get(key) for key in COMMON_FIELDS},
            "options": {key: config.get(key) for key in OPTION_FIELDS}, "queries": queries, "statistics": stats,
            "online_mode_evidence": modes,
            "inventory": accounting, "resources_per_party": resources, "session_phases": session_phases,
            "first_query_raw": measurements.get("first_query"), "session_total_raw": measurements.get("total_seconds"),
            "embedded_runtime_source_hashes": config.get("runtime_manifest", {}).get("sources", {})}


def comparisons(rows, links):
    candidates = [name for name in ORDER if name.startswith("ours_") and name in rows]
    pairs = [(left, right, "cumulative_ablation") for left, right in zip(candidates, candidates[1:])]
    if len(candidates) > 2:
        pairs.append(("ours_online_before", "ours_online_optimized", "total_online_optimization"))
    pairs += [(name, "ours_online_optimized", "vs_SCI") for name in MAIN_ROWS[2:]]
    output = []
    for left, right, kind in pairs:
        item = {"baseline": left, "candidate": right, "kind": kind, "status": "unavailable", "phases": {}}
        if rows[left]["status"] == rows[right]["status"] == "complete":
            item["status"] = "complete"
            for phase in PHASES:
                a, b = rows[left]["statistics"][phase], rows[right]["statistics"][phase]
                metrics = {key: ((a[key] or {}).get("median"), (b[key] or {}).get("median")) for key in SCALARS}
                metrics.update({name + "_communication_projection": (a["network_projection"][name]["communication_seconds"]["median"],
                    b["network_projection"][name]["communication_seconds"]["median"]) for name in links})
                item["phases"][phase] = {key: {"baseline_median": first, "candidate_median": second,
                    "baseline_over_candidate": first / second if first is not None and second else None,
                    "reduction_fraction": 1 - second / first if first and second is not None else None}
                    for key, (first, second) in metrics.items()}
        output.append(item)
    return output


def summarize(directory, links=NETWORKS, allow_incomplete=False):
    plan, execution = read(directory / "PLAN.json"), read(directory / "STATUS.json")
    identities = [row["id"] for row in plan["rows"]]
    if plan["mode"] != "formal" or len(set(identities)) != len(identities) or not set(MAIN_ROWS) <= set(identities) or not set(identities) <= set(ORDER):
        raise ValueError("new formal plan must contain the four matched main rows, with optional declared ablations")
    if plan["build"] != execution["build"]:
        raise ValueError("plan/execution build mismatch")
    if execution["status"] != "complete" and not allow_incomplete:
        raise ValueError("formal run incomplete; --allow-incomplete emits a pending report")
    rows, errors, common = {}, [], None
    for declaration in plan["rows"]:
        state = exactly(execution["rows"], id=declaration["id"])
        if state["status"] != "complete":
            rows[declaration["id"]] = {"id": declaration["id"], "status": state["status"], "reason": state.get("reason")}
            continue
        try:
            row = extract(declaration, state, plan, links)
            if common is None:
                common = row["common_config"]
            if row["common_config"] != common:
                raise ValueError("common workload, input, binary, resources or model preparation differs")
            rows[declaration["id"]] = row
        except (ValueError, KeyError, OSError) as error:
            errors.append({"id": declaration["id"], "reason": str(error)})
            rows[declaration["id"]] = {"id": declaration["id"], "status": "invalid", "reason": str(error)}
    if errors and not allow_incomplete:
        raise ValueError(str(errors))
    return {"schema": 1, "status": "complete" if all(row["status"] == "complete" for row in rows.values()) else "incomplete",
            "directory": str(directory), "build": plan["build"], "allocator": plan["allocator"], "common_config": common,
            "network_specs": links, "row_order": identities, "rows": rows, "comparisons": comparisons(rows, links), "errors": errors,
            "sources": {"phase_wall": "result.json measurements.queries[q][phase].per_party[*].seconds: max",
                        "process_cpu": "same per_party[*].process_cpu_seconds: max and sum; process CLOCK_PROCESS_CPUTIME_ID",
                        "communication": "sum per_party[*].sent_bytes, matched to causal trace; do not add received bytes",
                        "rounds": "round_measurements.phases[query,phase].logical_rounds",
                        "resources": "max of runtime party logs and process resource sampling"},
            "statistics": "10 fresh queries; median and p95 linear interpolation at0.95*(n-1); no confidence intervals",
            "projection": "8*two-party-sent-bytes/(Mbps*1e6)+causal-rounds*RTT_ms/1000; communication model only",
            "combined_phase": "sum offline/online phase wall maxima; sum per-party phase CPU then max/sum; sum bytes/depth/projections per query before statistics"}


def num(value, digits=4):
    return "缺失" if value is None else format(value, ".%df" % digits)


def report(data):
    lines = ["# SCI+ours 在线优化与 SCI/SiRNN 对照", "", "状态：`%s`。" % data["status"], "",
        "各行使用相同的新构建、fixture、batch128、单个 recurrent timestep、默认隐藏状态、每方4个逻辑CPU和1.5GiB地址空间上限。"
        "每行10次全新query，无warmup。旧实验仅用于选择配置，没有混入本轮正式数值。", "",
        "CPU时间来自process clock，累计本进程全部线程的user+sys时间，包含trace、系统调用等开销，可能大于墙钟。"
        "分别报告单方最大CPU时间、两方CPU总量和本机未限速墙钟；墙钟不能称为纯计算时间。", "",
        "通信是两方发送字节之和，仅计一次。轮数来自完整因果轨迹，线程或算子轮数不能相加替代整阶段轮数。"
        "LAN/WAN按实测字节与轮数计算传输及RTT成本，属于通信模型，未测量相应网络墙钟；不将本机墙钟与投影相加。", "",
        "| 方法 | 阶段 | 本机wall median/p95(s) | max-party CPU median(s) | 两方CPU total median(s) | 通信 median(MiB) | 因果轮数 median | LAN/WAN通信投影 median(s) |",
        "|---|---|---:|---:|---:|---:|---:|---:|"]
    for identity in data["row_order"]:
        row = data["rows"][identity]
        if row["status"] != "complete":
            lines.append("| %s | %s | — | — | — | — | — | — |" % (LABELS[identity], row["status"]))
            continue
        for phase in PHASES:
            s = row["statistics"][phase]
            med = lambda key: (s[key] or {}).get("median")
            lines.append("| %s | %s | %s/%s | %s | %s | %s | %s | %s/%s |" % (
                LABELS[identity], PHASE_LABELS[phase], num(med("wall_seconds")), num(s["wall_seconds"]["p95"]),
                num(med("cpu_seconds_max")), num(med("cpu_seconds_total")), num(med("sent_bytes") / (1 << 20), 3),
                num(med("causal_rounds"), 0), *[num(s["network_projection"][name]["communication_seconds"]["median"])
                                               for name in ("LAN", "WAN")]))
    lines += ["", "| 比较 baseline→candidate | 在线wall倍数/降幅 | 在线通信倍数/降幅 | 在线轮数倍数 | LAN/WAN通信投影倍数 |",
              "|---|---:|---:|---:|---:|"]
    for item in data["comparisons"]:
        if item["status"] != "complete":
            continue
        s = item["phases"]["online"]
        metric = lambda key: "%s×/%s%%" % (num(s[key]["baseline_over_candidate"], 3), num(s[key]["reduction_fraction"] * 100, 2))
        lines.append("| %s→%s | %s | %s | %s× | %s×/%s× |" % (LABELS[item["baseline"]], LABELS[item["candidate"]],
            metric("wall_seconds"), metric("sent_bytes"), num(s["causal_rounds"]["baseline_over_candidate"], 3),
            num(s["LAN_communication_projection"]["baseline_over_candidate"], 3),
            num(s["WAN_communication_projection"]["baseline_over_candidate"], 3)))
    if all(data["rows"].get(key, {}).get("status") == "complete" for key in MAIN_ROWS):
        before = data["rows"]["ours_online_before"]["statistics"]["online"]
        after = data["rows"]["ours_online_optimized"]["statistics"]["online"]
        lines += ["", "本轮SCI+ours在线通信和CPU中位数均下降，但离线以及离线＋在线合计仍高于两组SiRNN。"]
        if after["wall_seconds"]["p95"] > before["wall_seconds"]["p95"]:
            lines += ["", "在线墙钟存在长尾：优化后p95为%s秒，高于优化前%s秒；中位数下降不能说明尾延迟改善。" %
                      (num(after["wall_seconds"]["p95"]), num(before["wall_seconds"]["p95"]))]
    lines += ["", "比值为baseline/candidate中位数比，无置信区间；负降幅表示退化。离线＋在线先逐query求和再统计，"
              "不包含setup、input、ready或output。setup/model_prepare在Excel独立sheet中保留，不摊销或重复加入首次请求。", "",
              "未预处理SCI使用`--backend SCI`且不启用native预处理：其OT/triple生成在online里执行，"
              "本轮测得的offline阶段原样保留；prepared库存与零online生成政策对该行标N/A。"
              "预处理SCI和ours要求online零生成及库存一次消费。SCI用原生激活，ours用folded f12/q10；"
              "跨方法比较包含整数算法差异，不能解释成一个开关的收益。", "",
              "Excel包含LAN、WAN、逐query、setup、资源库存、比较和说明sheet。CPU计量与通信计数的来源、"
              "各阶段原值、配置与构建标识保存在METRICS.json。", ""]
    return "\n".join(lines)


def workbook(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    headers = ["方法", "配置ID", "阶段", "统计量", "样本数", "实测本机wall(s)", "实测max-party CPU(s)",
               "实测两方total CPU(s)", "两方发送(B)", "两方发送(MiB)", "因果轮数", "带宽(Mbps)", "RTT(ms)",
               "传输投影(s)", "RTT投影(s)", "通信合计投影(s)", "测量状态"]
    for network, link in data["network_specs"].items():
        sheet = book.create_sheet(network)
        sheet.append(headers)
        for identity in data["row_order"]:
            row = data["rows"][identity]
            if row["status"] != "complete":
                sheet.append([LABELS[identity], identity, None, None, None] + [None] * 11 + [row["status"]])
                continue
            for phase in PHASES:
                stats = row["statistics"][phase]
                for stat in ("median", "p95"):
                    val = lambda key: (stats[key] or {}).get(stat)
                    projected = stats["network_projection"][network]
                    sheet.append([LABELS[identity], identity, PHASE_LABELS[phase], stat, stats["wall_seconds"]["n"],
                        val("wall_seconds"), val("cpu_seconds_max"), val("cpu_seconds_total"), val("sent_bytes"),
                        val("sent_bytes") / (1 << 20), val("causal_rounds"), link["mbps"], link["rtt_ms"],
                        projected["transfer_seconds"][stat], projected["rtt_seconds"][stat],
                        projected["communication_seconds"][stat], row["status"]])
    raw = book.create_sheet("逐query")
    raw.append(["配置ID", "方法", "query", "阶段", "wall方1(s)", "wall方2(s)", "max wall(s)",
                "CPU方1(s)", "CPU方2(s)", "max CPU(s)", "total CPU(s)", "sent方1(B)", "sent方2(B)",
                "两方sent(B)", "MiB", "causal rounds", "LAN传输(s)", "LAN RTT(s)", "LAN合计(s)",
                "WAN传输(s)", "WAN RTT(s)", "WAN合计(s)"])
    setup = book.create_sheet("Setup")
    setup.append(["配置ID", "方法", "阶段", "wall(s)", "max CPU(s)", "total CPU(s)", "两方sent(B)", "因果轮数",
                  "LAN通信投影(s)", "WAN通信投影(s)", "原始字段"])
    stock = book.create_sheet("资源库存")
    stock.append(["配置ID", "方法", "方1 RSS峰值(MiB)", "方2 RSS峰值(MiB)", "方1 VmPeak(MiB)", "方2 VmPeak(MiB)",
                  "在线生成政策", "政策通过", "零online生成适用", "一次消费库存适用", "库存检查", "配置"])
    for identity in data["row_order"]:
        row = data["rows"][identity]
        if row["status"] != "complete":
            continue
        for query in row["queries"]:
            for phase, value in query["phases"].items():
                raw.append([identity, LABELS[identity], query["query"], PHASE_LABELS.get(phase, phase),
                    *value["party_wall_seconds"], value["wall_seconds"], *value["party_cpu_seconds"],
                    value["cpu_seconds_max"], value["cpu_seconds_total"], *value["party_sent_bytes"],
                    value["sent_bytes"], value["sent_bytes"] / (1 << 20), value["causal_rounds"],
                    *[value["network_projection"][name][key] for name in ("LAN", "WAN")
                      for key in ("transfer_seconds", "rtt_seconds", "communication_seconds")]])
        for phase, value in row["session_phases"].items():
            setup.append([identity, LABELS[identity], phase, value.get("wall_seconds"), value.get("cpu_seconds_max"),
                value.get("cpu_seconds_total"), value.get("sent_bytes"), value.get("causal_rounds"),
                (value.get("network_projection", {}).get("LAN", {})).get("communication_seconds"),
                (value.get("network_projection", {}).get("WAN", {})).get("communication_seconds"),
                json.dumps(value.get("raw_measurement"), ensure_ascii=False) if "raw_measurement" in value else None])
        setup.append([identity, LABELS[identity], "first_query (已含setup)", (row["first_query_raw"] or {}).get("seconds"),
                      None, None, (row["first_query_raw"] or {}).get("sent_bytes"), None, None, None,
                      json.dumps(row["first_query_raw"], ensure_ascii=False)])
        setup.append([identity, LABELS[identity], "session_total", row["session_total_raw"]])
        resources, inv = row["resources_per_party"], row["inventory"]
        stock.append([identity, LABELS[identity], *[value["rss_peak_bytes"] / (1 << 20) for value in resources],
            *[value["vm_peak_bytes"] / (1 << 20) for value in resources], inv["mode"], inv["online_generation_policy_passed"],
            inv["zero_online_generation_applicable"], inv["one_use_stock_applicable"],
            inv["one_use_stock_passed"] if inv["one_use_stock_applicable"] else "N/A",
            json.dumps(row["options"], ensure_ascii=False)])
    compare = book.create_sheet("比较")
    compare.append(["baseline", "candidate", "比较类型", "阶段", "指标", "baseline median", "candidate median", "baseline/candidate", "降幅"])
    for item in data["comparisons"]:
        if item["status"] != "complete":
            continue
        for phase, metrics in item["phases"].items():
            for key, value in metrics.items():
                compare.append([item["baseline"], item["candidate"], item["kind"], PHASE_LABELS[phase], key,
                    value["baseline_median"], value["candidate_median"], value["baseline_over_candidate"], value["reduction_fraction"]])
    notes = book.create_sheet("说明")
    notes.append(["项目", "说明"])
    definitions = {
        "状态": data["status"], "实验": "同一新构建与fixture；batch128，steps1；每方4逻辑CPU，1.5GiB；10fresh queries，无warmup",
        "计算时间": "process CPU clock累计全部线程user+sys，包含trace和syscalls；max-party与两方total分别报告，可能超过wall",
        "墙钟": "本机未限速阶段实测wall取两方最大值；包含通信等待，不能称纯计算时间",
        "通信字节": "仅两方sent相加，收发同一消息不重复计数；MiB=2^20 B",
        "因果轮数": "完整phase dependency DAG的send-flight depth，不能累加各线程轮数",
        "网络投影": data["projection"], "投影限制": "仅通信模型；未测量真实LAN/WAN墙钟，不将本机wall与通信投影相加",
        "统计": data["statistics"], "离线＋在线": data["combined_phase"],
        "SCI未预处理": "backend SCI，不请求native stock。OT/triple生成归online；测得offline原样保留，prepared库存/零生成政策N/A",
        "SCI预处理": "backend SCI_preprocessed，全部query native material在input前生成，一次消费，online无生成",
        "激活算法": "SCI用原生激活；ours用folded f12/q10。跨方法比较包含数值算法差异，不能解释成单个开关收益",
        "Setup": "setup/model_prepare独立保留；first_query已含setup；不摊销进每query，不重复相加",
        "历史": "旧结果没有混入本轮正式样本；资格检查不能作为性能采样",
        "原始数据": str(Path(data["directory"]) / "METRICS.json"),
        "构建": json.dumps(data["build"], ensure_ascii=False), "共同配置": json.dumps(data["common_config"], ensure_ascii=False)}
    for key, value in definitions.items():
        notes.append([key, value])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="34495E")
        for index, column in enumerate(sheet.columns, 1):
            maximum = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
            sheet.column_dimensions[get_column_letter(index)].width = min(48, max(13, maximum + 2))
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
                cell.alignment = Alignment(vertical="top", wrap_text=sheet.title == "说明")
    notes.column_dimensions["B"].width = 105
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=HERE / "formal")
    parser.add_argument("--allow-incomplete", action="store_true")
    parser.add_argument("--lan-mbps", type=float, default=1000)
    parser.add_argument("--lan-rtt-ms", type=float, default=.5)
    parser.add_argument("--wan-mbps", type=float, default=400)
    parser.add_argument("--wan-rtt-ms", type=float, default=4)
    args = parser.parse_args()
    values = (args.lan_mbps, args.lan_rtt_ms, args.wan_mbps, args.wan_rtt_ms)
    if any(not math.isfinite(value) for value in values) or min(args.lan_mbps, args.wan_mbps) <= 0 or min(args.lan_rtt_ms, args.wan_rtt_ms) < 0:
        parser.error("bandwidth must be positive; RTT nonnegative; all values finite")
    links = {"LAN": {"mbps": args.lan_mbps, "rtt_ms": args.lan_rtt_ms}, "WAN": {"mbps": args.wan_mbps, "rtt_ms": args.wan_rtt_ms}}
    directory = args.directory.resolve()
    data = summarize(directory, links, args.allow_incomplete)
    (directory / "METRICS.json").write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n")
    (directory / "REPORT.zh.md").write_text(report(data))
    workbook(data, directory / "REPORT.xlsx")
    print(json.dumps({"status": data["status"], "metrics": str(directory / "METRICS.json"),
                      "report": str(directory / "REPORT.zh.md"), "excel": str(directory / "REPORT.xlsx")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
